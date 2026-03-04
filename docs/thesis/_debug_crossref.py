"""Debug cross-references in the generated DOCX."""
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

doc_path = r"D:\KLTN\Quant wwebsite\quant-website\docs\thesis\generated\thesis_v2_draft.docx"
doc = Document(doc_path)

# Find paragraphs containing REF fields
count = 0
for para in doc.paragraphs:
    para_xml = etree.tostring(para._element, pretty_print=True).decode()
    if "instrText" in para_xml and "REF" in para_xml:
        count += 1
        if count <= 2:
            text = para.text
            print(f"=== Cross-ref paragraph {count} ===")
            print(f"Text: {text[:120]}")
            for elem in para._element.iter():
                if elem.tag == qn("w:instrText"):
                    print(f"  instrText: {elem.text}")
            print()

print(f"Total paragraphs with REF fields: {count}")
print()

# Check caption paragraphs with bookmarks
print("=== Caption paragraphs with bookmarks ===")
bm_count = 0
for para in doc.paragraphs:
    para_xml = etree.tostring(para._element, pretty_print=True).decode()
    if "bookmarkStart" in para_xml:
        for elem in para._element.iter():
            if elem.tag == qn("w:bookmarkStart"):
                n = elem.get(qn("w:name"))
                if n and n.startswith("_"):
                    bm_count += 1
                    print(f"  Caption: \"{para.text[:80]}\" -> bookmark: {n}")
                    break

print(f"Total caption bookmarks: {bm_count}")

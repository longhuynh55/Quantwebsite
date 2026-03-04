#!/usr/bin/env python3
"""
Fill school cover DOCX using python-docx with layout-safe defaults.

Default behavior:
- Replace existing cover markers only.
- Do NOT append extra pages.

Optional:
- Append acknowledgement/instructor page with --append-front-matter.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def load_data(path: Path) -> dict[str, str]:
    with path.open("r", encoding="utf-8") as f:
        raw = json.load(f)
    return {str(k): "" if v is None else str(v) for k, v in raw.items()}


def set_run_font(run, font_name: str = "Times New Roman", font_size_pt: float | None = None, bold: bool | None = None) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if font_size_pt is not None:
        run.font.size = Pt(font_size_pt)
    if bold is not None:
        run.bold = bold


def replace_in_runs(paragraph, old: str, new: str) -> int:
    count = 0
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            count += 1
    return count


def replace_in_story(story, old: str, new: str) -> int:
    replaced = 0
    for p in story.paragraphs:
        replaced += replace_in_runs(p, old, new)
    return replaced


def replace_global_tokens(doc: Document, data: dict[str, str]) -> None:
    # Main document
    for key, value in data.items():
        token = f"{{{{{key}}}}}"
        replace_in_story(doc, token, value)

    # Headers/footers
    for section in doc.sections:
        for header in section.header.paragraphs:
            for key, value in data.items():
                token = f"{{{{{key}}}}}"
                replace_in_runs(header, token, value)
        for footer in section.footer.paragraphs:
            for key, value in data.items():
                token = f"{{{{{key}}}}}"
                replace_in_runs(footer, token, value)


def apply_legacy_cover_mapping(doc: Document, data: dict[str, str]) -> None:
    thesis_title = data.get("THESIS_TITLE", "").strip()
    supervisor = data.get("SUPERVISOR_NAME", "").strip()
    student_name = data.get("STUDENT_NAME", "").strip()
    student_id = data.get("STUDENT_ID", "").strip()
    city = data.get("CITY", "").strip()
    year = data.get("YEAR", "").strip()

    for p in doc.paragraphs:
        replace_in_runs(p, "GVHD:", f"GVHD: {supervisor}" if supervisor else "GVHD:")
        replace_in_runs(p, "SVTH:", f"SVTH: {student_name}" if student_name else "SVTH:")
        replace_in_runs(p, "MSSV:", f"MSSV: {student_id}" if student_id else "MSSV:")

        if thesis_title:
            replace_in_runs(p, "TÊN ĐỀ TÀI", thesis_title)
            replace_in_runs(p, "TEN DE TAI", thesis_title)

        if city:
            replace_in_runs(p, "TP.HCM,", f"{city},")

        if year:
            replace_in_runs(p, "/….", year)
            replace_in_runs(p, "/...", year)


def configure_front_matter_page(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)


def add_paragraph(doc: Document, text: str, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, size=13, space_after=6, line_spacing=1.5) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    fmt = p.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing
    run = p.add_run(text)
    set_run_font(run, font_name="Times New Roman", font_size_pt=size, bold=bold)


def append_front_matter(doc: Document, data: dict[str, str]) -> None:
    ack_text = data.get("ACKNOWLEDGEMENT_TEXT", "").strip()
    instructor_comment = data.get("INSTRUCTOR_COMMENT", "").strip()
    if not ack_text and not instructor_comment:
        return

    title = data.get("ACKNOWLEDGEMENT_TITLE", "").strip() or "ACKNOWLEDGEMENTS"
    student_name = data.get("STUDENT_NAME", "").strip()

    doc.add_page_break()
    configure_front_matter_page(doc.sections[-1])

    add_paragraph(
        doc,
        title,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        size=16,
        space_after=12,
        line_spacing=1.2,
    )

    if ack_text:
        for line in ack_text.splitlines():
            if line.strip():
                add_paragraph(doc, line.strip(), align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, size=13, space_after=6)
            else:
                add_paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, size=13, space_after=6)

    if instructor_comment:
        add_paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, size=13, space_after=6)
        add_paragraph(doc, "Instructor Comment", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=13, space_after=6)
        for line in instructor_comment.splitlines():
            add_paragraph(doc, line.strip(), align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=13, space_after=6)

    if student_name:
        add_paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=13, space_after=6)
        add_paragraph(doc, f"Student: {student_name}", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=13, space_after=6)


def main() -> int:
    parser = argparse.ArgumentParser(description="Fill school cover DOCX with python-docx.")
    parser.add_argument("--template", required=True, help="Path to source cover DOCX.")
    parser.add_argument("--data", required=True, help="Path to JSON data.")
    parser.add_argument("--output", required=True, help="Path to output DOCX.")
    parser.add_argument(
        "--append-front-matter",
        action="store_true",
        help="Append acknowledgement/instructor page after cover.",
    )
    args = parser.parse_args()

    template = Path(args.template).resolve()
    data_path = Path(args.data).resolve()
    output = Path(args.output).resolve()

    if not template.exists():
        raise FileNotFoundError(f"Template not found: {template}")
    if not data_path.exists():
        raise FileNotFoundError(f"Data not found: {data_path}")

    data = load_data(data_path)
    doc = Document(str(template))

    replace_global_tokens(doc, data)
    apply_legacy_cover_mapping(doc, data)

    if args.append_front_matter:
        append_front_matter(doc, data)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    print(f"DOCX generated: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


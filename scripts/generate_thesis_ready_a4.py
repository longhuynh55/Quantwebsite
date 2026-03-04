#!/usr/bin/env python3
"""
Generate a ready-to-use thesis DOCX with strict A4 layout.

This script follows the old generate_thesis.py approach:
- python-docx based layout control
- A4 page size and thesis margins
- Times New Roman + 1.5 line spacing
- front matter pages before chapters
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def load_json(path: Path) -> dict[str, str]:
    with path.open("r", encoding="utf-8-sig") as f:
        raw = json.load(f)
    return {str(k): "" if v is None else str(v) for k, v in raw.items()}


def set_run_font(run, name: str = "Times New Roman", size: float = 13, bold: bool = False, italic: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_a4_layout(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)


def apply_cover_mapping(doc: Document, data: dict[str, str]) -> None:
    thesis_title = data.get("THESIS_TITLE", "").strip()
    supervisor = data.get("SUPERVISOR_NAME", "").strip()
    student = data.get("STUDENT_NAME", "").strip()
    student_id = data.get("STUDENT_ID", "").strip()
    student_class = data.get("CLASS", "").strip()
    city = data.get("CITY", "").strip() or "Ho Chi Minh City"
    year = data.get("YEAR", "").strip()

    for p in doc.paragraphs:
        txt = p.text
        if "GVHD:" in txt:
            p.text = txt.replace("GVHD:", f"GVHD: {supervisor}" if supervisor else "GVHD:")
        if "SVTH:" in p.text:
            p.text = p.text.replace("SVTH:", f"SVTH: {student}" if student else "SVTH:")
        if "MSSV:" in p.text:
            p.text = p.text.replace("MSSV:", f"MSSV: {student_id}" if student_id else "MSSV:")
        if "Class" in p.text and ":" in p.text and student_class:
            p.text = p.text.replace("Class", "Class").replace(":", f": {student_class}", 1)
        if thesis_title:
            p.text = p.text.replace("TÃŠN Äá»€ TÃ€I", thesis_title)
            p.text = p.text.replace("TEN DE TAI", thesis_title)
        if city and "TP.HCM," in p.text:
            p.text = p.text.replace("TP.HCM,", f"{city},")
        if year:
            p.text = p.text.replace("/â€¦.", year)
            p.text = p.text.replace("/...", year)

        for r in p.runs:
            set_run_font(r, size=13)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(18 if level == 1 else 12)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.5
    pf.first_line_indent = None
    run = p.add_run(text)
    size = 16 if level == 1 else (14 if level == 2 else 13)
    set_run_font(run, size=size, bold=True)


def add_body(doc: Document, text: str, align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent_cm: float = 1.27) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(first_line_indent_cm) if first_line_indent_cm > 0 else None
    run = p.add_run(text)
    set_run_font(run, size=13)


def add_center_line(doc: Document, text: str, size: float = 13, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.5
    pf.first_line_indent = None
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def add_empty_lines(doc: Document, n: int) -> None:
    for _ in range(n):
        add_center_line(doc, "", size=13, bold=False)


def create_generated_cover(doc: Document, data: dict[str, str]) -> None:
    title = data.get("THESIS_TITLE", "").strip() or "THESIS TITLE"
    student = data.get("STUDENT_NAME", "").strip() or "Student Name"
    student_id = data.get("STUDENT_ID", "").strip() or "Student ID"
    supervisor = data.get("SUPERVISOR_NAME", "").strip() or "Supervisor"
    city = data.get("CITY", "").strip() or "Ho Chi Minh City"
    year = data.get("YEAR", "").strip() or "2026"
    faculty = data.get("FACULTY_NAME", "").strip() or "FACULTY OF FINANCE AND BANKING"

    add_center_line(doc, "VIET NAM NATIONAL UNIVERSITY HO CHI MINH CITY", size=13, bold=True)
    add_center_line(doc, "UNIVERSITY OF ECONOMICS AND LAW", size=14, bold=True)
    add_center_line(doc, faculty.upper(), size=14, bold=True)
    add_empty_lines(doc, 2)
    add_center_line(doc, "GRADUATION THESIS", size=16, bold=True)
    add_empty_lines(doc, 1)

    words = title.split()
    line = []
    wrapped: list[str] = []
    for w in words:
        test = " ".join(line + [w]).strip()
        if len(test) > 46 and line:
            wrapped.append(" ".join(line))
            line = [w]
        else:
            line.append(w)
    if line:
        wrapped.append(" ".join(line))
    for t in wrapped:
        add_center_line(doc, t.upper(), size=16, bold=True)

    add_empty_lines(doc, 3)
    add_center_line(doc, f"Supervisor: {supervisor}", size=13, bold=True)
    add_center_line(doc, f"Student name: {student}", size=13, bold=True)
    add_center_line(doc, f"Student ID: {student_id}", size=13, bold=True)
    add_empty_lines(doc, 4)
    add_center_line(doc, f"{city}, {year}", size=13, bold=True)


def add_front_matter_pages(doc: Document, data: dict[str, str]) -> None:
    ack_title = (data.get("ACKNOWLEDGEMENT_TITLE", "") or "ACKNOWLEDGEMENTS").strip()
    ack_text = data.get("ACKNOWLEDGEMENT_TEXT", "").strip()

    # i - Acknowledgements
    doc.add_page_break()
    add_center_line(doc, "i")
    add_center_line(doc, ack_title, size=15, bold=True)
    if ack_text:
        for line in ack_text.splitlines():
            if line.strip():
                add_body(doc, line.strip(), first_line_indent_cm=1.27)
            else:
                add_body(doc, "", first_line_indent_cm=0)
    else:
        add_body(doc, "Acknowledgement content will be finalized before submission.", first_line_indent_cm=1.27)

    # ii - Instructor's comment
    doc.add_page_break()
    add_center_line(doc, "ii")
    add_center_line(doc, "INSTRUCTOR'S COMMENT", size=15, bold=True)
    for _ in range(10):
        add_body(doc, "." * 115, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent_cm=0)

    # iii - TOC
    doc.add_page_break()
    add_center_line(doc, "iii")
    add_center_line(doc, "TABLE OF CONTENTS", size=15, bold=True)
    add_body(doc, "Generate this page automatically in Word after finalizing heading styles.", first_line_indent_cm=0)

    # iv - List of tables and figures
    doc.add_page_break()
    add_center_line(doc, "iv")
    add_center_line(doc, "LIST OF TABLES, FIGURES", size=15, bold=True)
    add_body(doc, "Generate this page automatically from captions.", first_line_indent_cm=0)

    # v - Abbreviations
    doc.add_page_break()
    add_center_line(doc, "v")
    add_center_line(doc, "LIST OF ABBREVIATIONS", size=15, bold=True)
    for row in [
        "API - Application Programming Interface",
        "ML - Machine Learning",
        "LLM - Large Language Model",
        "HOSE - Ho Chi Minh City Stock Exchange",
    ]:
        add_body(doc, row, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent_cm=0)

    # vi - Abstract
    doc.add_page_break()
    add_center_line(doc, "vi")
    add_center_line(doc, "ABSTRACT", size=15, bold=True)
    add_body(
        doc,
        "This abstract section is a placeholder and should be replaced with a final 200-300 word abstract.",
        first_line_indent_cm=1.27,
    )


def add_prechapter_pages(doc: Document) -> None:
    doc.add_page_break()
    add_center_line(doc, "TABLE OF CONTENTS", size=15, bold=True)
    add_body(doc, "Generate this page automatically in Word after finalizing heading styles.", first_line_indent_cm=0)

    doc.add_page_break()
    add_center_line(doc, "LIST OF TABLES, FIGURES", size=15, bold=True)
    add_body(doc, "Generate this page automatically from captions.", first_line_indent_cm=0)

    doc.add_page_break()
    add_center_line(doc, "LIST OF ABBREVIATIONS", size=15, bold=True)
    for row in [
        "API - Application Programming Interface",
        "ML - Machine Learning",
        "LLM - Large Language Model",
        "HOSE - Ho Chi Minh City Stock Exchange",
    ]:
        add_body(doc, row, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent_cm=0)

    doc.add_page_break()
    add_center_line(doc, "ABSTRACT", size=15, bold=True)
    add_body(
        doc,
        "This abstract section is a placeholder and should be replaced with a final 200-300 word abstract.",
        first_line_indent_cm=1.27,
    )


def markdown_lines(path: Path) -> list[str]:
    return path.read_text(encoding="utf-8").splitlines()


def append_markdown(doc: Document, md_path: Path) -> None:
    doc.add_page_break()
    for line in markdown_lines(md_path):
        s = line.rstrip()
        if s.startswith("### "):
            add_heading(doc, s[4:].strip(), level=3)
            continue
        if s.startswith("## "):
            add_heading(doc, s[3:].strip(), level=2)
            continue
        if s.startswith("# "):
            add_heading(doc, s[2:].strip(), level=1)
            continue
        if s.startswith("- "):
            add_body(doc, f"- {s[2:].strip()}", align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent_cm=0)
            continue
        if not s:
            add_body(doc, "", first_line_indent_cm=0)
            continue
        add_body(doc, s)


def append_references(doc: Document, refs_path: Path) -> None:
    doc.add_page_break()
    add_heading(doc, "References", level=1)
    for line in markdown_lines(refs_path):
        s = line.strip()
        if not s or s.startswith("#"):
            continue
        add_body(doc, s, first_line_indent_cm=0)


def save_combined_markdown(chapters: Iterable[Path], refs: Path, output: Path) -> None:
    blocks: list[str] = []
    for p in chapters:
        blocks.append(f"<!-- BEGIN: {p.name} -->")
        blocks.append(p.read_text(encoding="utf-8"))
        blocks.append(f"<!-- END: {p.name} -->")
        blocks.append("")
    blocks.append("<!-- BEGIN: references_apa.md -->")
    blocks.append(refs.read_text(encoding="utf-8"))
    blocks.append("<!-- END: references_apa.md -->")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n".join(blocks), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate strict-A4 ready thesis DOCX from cover template + markdown chapters.")
    parser.add_argument("--template", required=False, default="", help="Path to base cover DOCX template (used when --cover-mode template).")
    parser.add_argument("--cover-mode", choices=["generated", "template"], default="generated", help="Choose generated cover page or template-based cover.")
    parser.add_argument("--skip-cover-mapping", action="store_true", help="Do not replace cover fields in template mode.")
    parser.add_argument("--skip-front-matter", action="store_true", help="Do not auto-generate acknowledgement/comment/TOC pages.")
    parser.add_argument("--add-prechapter-pages", action="store_true", help="Insert TOC/List/Abbreviations/Abstract pages before Chapter 1.")
    parser.add_argument("--cover-data", required=True, help="Path to JSON data for cover and front matter.")
    parser.add_argument("--chapters", nargs="+", required=True, help="Markdown chapter files in desired order.")
    parser.add_argument("--references", required=True, help="APA references markdown file.")
    parser.add_argument("--output-docx", required=True, help="Output DOCX path.")
    parser.add_argument("--output-markdown", required=True, help="Output combined markdown path.")
    args = parser.parse_args()

    cover_data_path = Path(args.cover_data).resolve()
    chapters = [Path(p).resolve() for p in args.chapters]
    references = Path(args.references).resolve()
    output_docx = Path(args.output_docx).resolve()
    output_md = Path(args.output_markdown).resolve()

    if not cover_data_path.exists():
        raise FileNotFoundError(f"Cover data not found: {cover_data_path}")
    for ch in chapters:
        if not ch.exists():
            raise FileNotFoundError(f"Chapter file not found: {ch}")
    if not references.exists():
        raise FileNotFoundError(f"References file not found: {references}")

    data = load_json(cover_data_path)
    if args.cover_mode == "template":
        if not args.template:
            raise ValueError("Template path is required when --cover-mode template.")
        template = Path(args.template).resolve()
        if not template.exists():
            raise FileNotFoundError(f"Template not found: {template}")
        doc = Document(str(template))
        set_a4_layout(doc)
        if not args.skip_cover_mapping:
            apply_cover_mapping(doc, data)
    else:
        doc = Document()
        set_a4_layout(doc)
        create_generated_cover(doc, data)

    if not args.skip_front_matter:
        add_front_matter_pages(doc, data)
    if args.add_prechapter_pages:
        add_prechapter_pages(doc)
    for ch in chapters:
        append_markdown(doc, ch)
    append_references(doc, references)
    set_a4_layout(doc)

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_docx))
    save_combined_markdown(chapters, references, output_md)

    print(f"Ready DOCX: {output_docx}")
    print(f"Combined markdown: {output_md}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
